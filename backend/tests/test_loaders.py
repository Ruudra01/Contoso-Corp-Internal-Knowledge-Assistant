"""One test per supported document type, plus the routing and failure paths.

Each format must reduce to the same `Block` representation, so the assertions are
shared rather than written per format.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.core.errors import DocumentParseError, UnsupportedFormatError
from app.ingestion.loaders import (
    DocxLoader,
    HtmlLoader,
    MarkdownLoader,
    PdfLoader,
    loader_for,
    supported_extensions,
)
from app.ingestion.models import BlockKind, DocumentType

ALL_TYPES = ["pdf", "docx", "html", "markdown"]


@pytest.mark.parametrize("sample", ALL_TYPES, indirect=True)
def test_loader_produces_common_representation(sample: Path) -> None:
    raw = loader_for(sample).load(sample)

    assert raw.blocks, "loader produced no blocks"
    assert isinstance(raw.document_type, DocumentType)
    assert raw.source_path == str(sample)
    # Every block is a valid member of the shared vocabulary with non-empty text.
    for block in raw.blocks:
        assert isinstance(block.kind, BlockKind)
        assert block.text.strip()
        assert block.level is None or 1 <= block.level <= 6
    # Structure was recovered, not just a text dump.
    assert any(b.is_heading for b in raw.blocks)
    assert raw.title_hint


@pytest.mark.parametrize(
    ("sample", "expected"),
    [
        ("pdf", DocumentType.PDF),
        ("docx", DocumentType.DOCX),
        ("html", DocumentType.HTML),
        ("markdown", DocumentType.MARKDOWN),
    ],
    indirect=["sample"],
)
def test_loader_reports_its_document_type(sample: Path, expected: DocumentType) -> None:
    assert loader_for(sample).load(sample).document_type is expected


def test_registry_routes_every_supported_extension() -> None:
    assert set(supported_extensions()) >= {".pdf", ".docx", ".html", ".md"}


def test_registry_rejects_unknown_extension() -> None:
    with pytest.raises(UnsupportedFormatError, match="no loader"):
        loader_for(Path("policy.rtf"))


# -- format-specific guarantees ------------------------------------------------


def test_pdf_preserves_page_numbers(corpus: Path) -> None:
    raw = PdfLoader().load(corpus / "pdf" / "cnt-hr-005_paid_time_off_policy.pdf")

    assert raw.page_count == 1
    assert all(b.page_number == 1 for b in raw.blocks)


def test_pdf_infers_title_heading_from_typography(corpus: Path) -> None:
    """These PDFs carry no outline, so the 20pt bold title must be recovered
    from font size alone."""
    raw = PdfLoader().load(corpus / "pdf" / "cnt-hr-005_paid_time_off_policy.pdf")
    headings = [b for b in raw.blocks if b.is_heading]

    assert headings[0].level == 1
    assert headings[0].text == "Paid Time Off Policy"


def test_pdf_joins_wrapped_lines_into_one_paragraph(corpus: Path) -> None:
    """A sentence wrapped across two lines must not become two blocks."""
    raw = PdfLoader().load(corpus / "pdf" / "cnt-hr-005_paid_time_off_policy.pdf")
    accrual = next(b for b in raw.blocks if b.text.startswith("Accrual:"))

    assert "25 PTO days annually beginning with their sixth year." in accrual.text


def test_pdf_rejects_file_without_text_layer(tmp_path: Path) -> None:
    import pymupdf

    blank = tmp_path / "scan.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(str(blank))
    document.close()

    with pytest.raises(DocumentParseError, match="no text layer"):
        PdfLoader().load(blank)


def test_docx_maps_styles_to_heading_levels(corpus: Path) -> None:
    raw = DocxLoader().load(corpus / "docx" / "cnt-fin-016_expense_reimbursement_procedure.docx")
    headings = [b for b in raw.blocks if b.is_heading]

    assert headings[0].level == 1
    assert headings[0].text == "Expense Reimbursement Procedure"


def test_docx_keeps_body_order_and_renders_tables(tmp_path: Path) -> None:
    """`document.paragraphs` and `document.tables` are separate sequences, so
    order has to come from the body XML."""
    import docx

    document = docx.Document()
    document.add_heading("Approvals", level=1)
    document.add_paragraph("Before the table.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Request"
    table.cell(0, 1).text = "Approver"
    table.cell(1, 0).text = "PTO"
    table.cell(1, 1).text = "Manager"
    document.add_paragraph("After the table.")
    path = tmp_path / "ordered.docx"
    document.save(str(path))

    kinds = [b.kind for b in DocxLoader().load(path).blocks]
    assert kinds == [
        BlockKind.HEADING,
        BlockKind.PARAGRAPH,
        BlockKind.TABLE,
        BlockKind.PARAGRAPH,
    ]


def test_docx_rejects_non_docx_file(tmp_path: Path) -> None:
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip archive")

    with pytest.raises(DocumentParseError, match="cannot read docx"):
        DocxLoader().load(broken)


def test_html_maps_h1_h6_to_levels_and_drops_chrome(tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>T</title><style>p{color:red}</style></head><body>"
        "<nav>Home | Away</nav>"
        "<h1>Policy</h1><h2>Scope</h2><p>Applies to staff.</p>"
        "<ul><li>First</li></ul>"
        "<script>alert(1)</script><footer>(c) Contoso</footer></body></html>",
        encoding="utf-8",
    )

    blocks = HtmlLoader().load(path).blocks
    text = " ".join(b.text for b in blocks)

    assert [(b.kind, b.level) for b in blocks] == [
        (BlockKind.HEADING, 1),
        (BlockKind.HEADING, 2),
        (BlockKind.PARAGRAPH, None),
        (BlockKind.LIST_ITEM, None),
    ]
    assert "alert" not in text and "Home | Away" not in text and "(c) Contoso" not in text


def test_html_renders_table_as_pipe_table(tmp_path: Path) -> None:
    path = tmp_path / "table.html"
    path.write_text(
        "<body><h1>M</h1><table><tr><th>A</th><th>B</th></tr>"
        "<tr><td>1</td><td>2</td></tr></table></body>",
        encoding="utf-8",
    )

    table = next(b for b in HtmlLoader().load(path).blocks if b.kind is BlockKind.TABLE)

    assert table.text.splitlines()[0] == "| A | B |"
    assert table.text.splitlines()[2] == "| 1 | 2 |"
    assert table.is_atomic


def test_html_rejects_empty_body(tmp_path: Path) -> None:
    path = tmp_path / "empty.html"
    path.write_text("<html><body></body></html>", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="no content"):
        HtmlLoader().load(path)


def test_markdown_maps_atx_headings_to_levels(corpus: Path) -> None:
    raw = MarkdownLoader().load(corpus / "markdown" / "cnt-it-024_it_equipment_asset_policy.md")
    levels = [(b.level, b.text) for b in raw.blocks if b.is_heading]

    assert (1, "IT Equipment & Asset Policy") in levels
    assert (2, "Company equipment") in levels
    assert (2, "Return") in levels


def test_markdown_keeps_table_intact(corpus: Path) -> None:
    raw = MarkdownLoader().load(corpus / "markdown" / "cnt-hr-021_manager_approval_matrix.md")
    table = next(b for b in raw.blocks if b.kind is BlockKind.TABLE)
    lines = table.text.splitlines()

    assert lines[0].startswith("| Request | Manager |")
    # Header, separator, and one row per matrix entry.
    assert len(lines) == 9
    assert "Five-day remote schedule" in table.text


def test_markdown_strips_emphasis_markers(tmp_path: Path) -> None:
    path = tmp_path / "doc.md"
    path.write_text("# T\n\nRequests of **three or fewer** days.\n", encoding="utf-8")

    body = next(b for b in MarkdownLoader().load(path).blocks if not b.is_heading)

    assert body.text == "Requests of three or fewer days."


def test_markdown_rejects_empty_file(tmp_path: Path) -> None:
    path = tmp_path / "empty.md"
    path.write_text("   \n", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="no content"):
        MarkdownLoader().load(path)
